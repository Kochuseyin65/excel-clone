"""Asus ROG Strix + RTX 4060 + Intel Core i7 13. nesil ilan tablosunu (sadece Türkiye, xlsx + csv) üretir.

Veriler 12 Ekim 2026'da web aramasıyla toplandı. Fiyatlar ve stok durumu
anlık görüntüdür; satın almadan önce linkten kontrol edin.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).parent
BASENAME = "asus_rog_strix_rtx4060_i7_13nesil_turkiye"

HEADERS = [
    "No", "Durum", "Seri", "Model Kodu", "İşlemci", "Ekran Kartı", "RAM",
    "Depolama (SSD)", "Ekran", "İşletim Sistemi", "Fiyat", "Para Birimi",
    "Satıcı / Platform", "Stok / İlan Durumu",
    "Kondisyon / Açıklama (ikinci el için)", "Veri Tarihi", "Link",
]

CPU = "Intel Core i7-13650HX (13. nesil)"
GPU = "RTX 4060 8GB"
G16_SCREEN = '16" WUXGA 1920x1200 165Hz'
G18_SCREEN = '18" QHD+ 2560x1600 240Hz'

# (durum, seri, model kodu, ram, ssd, ekran, os, fiyat, para birimi, satıcı, ülke, stok, kondisyon, tarih, link)
ROWS = [
    # ---------------- SIFIR ----------------
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3110W", "16GB DDR5", "512GB", G16_SCREEN, "Windows 11", 49366, "TL",
     "Vatan Bilgisayar", "TR", "Tükendi", "", "Ekim 2026",
     "https://www.vatanbilgisayar.com/asus-rog-strix-g16-13-nesil-core-i7-13650hx-16gb-512gb-16inc-rtx4060-8gb-w11.html"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3110W", "16GB DDR5", "512GB", G16_SCREEN, "Windows 11 Home", 61930.80, "TL",
     "Cimri (en ucuz)", "TR", "Son fiyat", "", "06.04.2026",
     "https://www.cimri.com/oyun-bilgisayari/en-ucuz-asus-rog-strix-g16-g614jv-n3110w-i7-13650hx-16gb-ram-512gb-ssd-rtx-4060-windows-11-home-16-inc-oyun-bilgisayari-fiyatlari,2342840887"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3110", "16GB DDR5", "512GB", G16_SCREEN, "FreeDOS", 89999, "TL",
     "Hepsiburada", "TR", "Listede", "GPU 140W", "Ekim 2026",
     "https://www.hepsiburada.com/asus-rog-strix-g16-g614jv-n3110-intel-core-i7-13650hx-16gb-512gb-ssd-rtx4060-freedos-16-wuxga-165hz-tasinabilir-bilgisayar-pm-HBC00003X9GTP"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3110", "16GB DDR5", "512GB", G16_SCREEN, "FreeDOS", 64599, "TL",
     "Cimri (en ucuz)", "TR", "Son fiyat", "", "29.12.2025",
     "https://www.cimri.com/oyun-bilgisayari/en-ucuz-asus-rog-strix-g16-g614jv-n3110-i7-13650hx-16gb-ram-512gb-ssd-rtx-4060-freedos-16-inc-oyun-bilgisayari-fiyatlari,2346829360"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3242W", "16GB DDR5", "512GB", G16_SCREEN, "Windows 11", 79999, "TL",
     "Hepsiburada", "TR", "Listede", "GPU maks. 175W TGP, MUX switch", "Ekim 2026",
     "https://www.hepsiburada.com/asus-rog-strix-g16-g614jv-n3242w-intel-core-i7-13650hx-16-gb-512-gb-ssd-rtx4060-windows-11-16-wuxga-165hz-tasinabilir-bilgisayar-pm-HBC000068UP9G"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3242W", "16GB DDR5", "512GB", G16_SCREEN, "Windows 11", 69997, "TL",
     "Teknosa", "TR", "Stokta yok", "", "Ekim 2026",
     "https://www.teknosa.com/asus-rog-strix-g16-g614jvn3242w-intel-core-i713650hx-16gb-ram-512gb-ssd-geforce-rtx-4060-4gb-16-w11-gri-gaming-notebook-p-125035739"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3157W", "16GB DDR5", "512GB", G16_SCREEN, "Windows 11", 71278, "TL",
     "MediaMarkt (OnuAl üzerinden)", "TR", "Stok bilinmiyor", "Fiyat geçmişi 54.997 – 85.499 TL", "Temmuz 2025",
     "https://onual.com/urun/asus-rog-strix-g16-g614jv-n3157w-core-i7-13650hx-islemci-16gb-ram-512gb-ssd-rtx-4060-ekran-karti-16-p-1339258"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3157W04", "32GB DDR5 (yükseltilmiş)", "512GB", G16_SCREEN, "Windows 11", 88499, "TL",
     "Hepsiburada", "TR", "Listede", "GPU 140W; RAM satıcı tarafından yükseltilmiş", "Ekim 2026",
     "https://www.hepsiburada.com/asus-rog-strix-g16-g614jv-n3157w-gaming-core-i7-13650hx-ddr5-islemci-16gb-ram-512gb-ssd-rtx4060-ekran-karti-16-win11-gaming-laptop-gri-pm-HBC00005C7ELB"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3157", "16GB DDR5", "512GB", G16_SCREEN, "FreeDOS", None, "TL",
     "Teknosa", "TR", "Stokta yok", "", "Ekim 2026",
     "https://www.teknosa.com/asus-rog-strix-g16-g614jvn3157-intel-core-i7-13650hx-16-16-gb-ram-512-gb-ssd-8-gb-rtx-4060-fhd-freedos-gaming-laptop-p-785354233"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3158", "16GB DDR5", "512GB", G16_SCREEN, "FreeDOS", None, "TL",
     "Hepsiburada", "TR", "Fiyat görülemedi", "", "Ekim 2026",
     "https://www.hepsiburada.com/asus-rog-strix-g16-g614jv-n3158-intel-i7-13650hx-16gb-ddr5-512gb-pcie-ssd-8gb-rtx4060-16-wuxga-165hz-freedos-pm-HBC00003XDZ83"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3452-A73", "24GB DDR5 (yükseltilmiş)", "512GB", G16_SCREEN, "Windows 11 Pro", 99999, "TL",
     "Hepsiburada", "TR", "Listede", "RAM satıcı tarafından yükseltilmiş", "Ekim 2026",
     "https://www.hepsiburada.com/asus-rog-strix-g16-intel-core-i7-13650hx-24gb-512gb-ssd-rtx4060-8gb-165hz-windows-11-pro-16-wuxga-tasinabilir-bilgisayar-g614jv-n3452-a73-pm-HBC0000726E0A"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3452-A78", "32GB DDR5 (yükseltilmiş)", "1TB", G16_SCREEN, "Windows 11 Pro", 155346.38, "TL",
     "Akakçe (4 satıcı, en ucuz)", "TR", "Stokta", "RAM/SSD satıcı tarafından yükseltilmiş", "07.10.2026",
     "https://www.akakce.com/oyun-bilgisayari/en-ucuz-asus-rog-strix-g16-intel-core-i7-13650hx-32gb-1tb-ssd-rtx4060-8gb-165hz-windows-11-pro-16-wuxga-tasinabilir-bilgisayar-g614jv-n3452-a78-fiyati,1102126746.html"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3452-A32", "96GB DDR5 (yükseltilmiş)", "4TB", G16_SCREEN, "FreeDOS", 189999, "TL",
     "Hepsiburada", "TR", "Listede", "RAM/SSD satıcı tarafından yükseltilmiş", "Ekim 2026",
     "https://www.hepsiburada.com/asus-rog-strix-g16-intel-core-i7-13650hx-96gb-4tb-ssd-rtx4060-8gb-165hz-freedos-16-wuxga-tasinabilir-bilgisayar-g614jv-n3452-a32-pm-HBC0000726E3B"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3452", "32GB DDR5", "512GB", G16_SCREEN, "FreeDOS", None, "TL",
     "Turkcell Pasaj / Cimri", "TR", "Stokta yok", "", "Ekim 2026",
     "https://www.turkcell.com.tr/pasaj/bilgisayar-tablet/bilgisayarlar/laptoplar/asus-rog-strix-g16-g614jv-n3452-i7-13650hx-32-gb-512-gb-ssd-rtx4060-16-inc-free-dos-dizustu-bilgisayar"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3456", "16GB DDR5", "512GB", G16_SCREEN, "-", 49999, "TL",
     "Akakçe", "TR", "Güncel fiyat yok", "Son görülen fiyat 14.10.2024", "14.10.2024",
     "https://www.akakce.com/oyun-bilgisayari/en-ucuz-asus-rog-strix-g16-g614jv-n3456-i7-13650hx-16-gb-512-gb-ssd-rtx4060-16-wuxga-gaming-laptop-fiyati,717631126.html"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JVA26-N3159A26", "16GB DDR5", "1TB (yükseltilmiş)", '16" FHD+ 165Hz', "FreeDOS", 51360.90, "TL",
     "Akakçe (Teknosa ~51.845 TL)", "TR", "Stokta", "", "Ekim 2026",
     "https://www.akakce.com/laptop-notebook/en-ucuz-asus-rog-strix-g16-g614jva26-n3159a26-intel-core-i7-13650hx-16gb-1tb-ssd-rtx-4060-freedos-16-fhd-165hz-tasinabilir-bilgisayar-fiyati,497473709.html"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JVA44-N3159A44", "32GB DDR5 (yükseltilmiş)", "4TB (yükseltilmiş)", '16" FHD+ 165Hz', "FreeDOS", 61011, "TL",
     "Akakçe", "TR", "Stokta", "Ücretsiz kargo", "Ekim 2026",
     "https://akakce.com/laptop-notebook/en-ucuz-asus-rog-strix-g16-g614jva44-n3159a44-intel-core-i7-13650hx-32gb-4tb-ssd-rtx-4060-freedos-16-fhd-165hz-tasinabilir-bilgisayar-fiyati,497473419.html"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JVA10-N3159A10", "8GB DDR5", "1TB (yükseltilmiş)", '16" FHD+ 165Hz', "FreeDOS", None, "TL",
     "Akakçe", "TR", "Fiyat görülemedi", "", "Ekim 2026",
     "https://www.akakce.com/laptop-notebook/en-ucuz-asus-rog-strix-g16-g614jva10-n3159a10-intel-core-i7-13650hx-8gb-1tb-ssd-rtx-4060-freedos-16-fhd-165hz-tasinabilir-bilgisayar-fiyati,497512908.html"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JVA42-N3159A42", "32GB DDR5 (yükseltilmiş)", "1TB (yükseltilmiş)", '16" FHD+ 165Hz', "FreeDOS", None, "TL",
     "Hepsiburada", "TR", "Fiyat görülemedi", "", "Ekim 2026",
     "https://www.hepsiburada.com/asus-rog-strix-g16-g614jva42-n3159a42-intel-core-i7-13650hx-32gb-1tb-ssd-rtx-4060-freedos-16-fhd-165hz-tasinabilir-bilgisayar-pm-HBC000046KWYX"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3159A14", "64GB DDR5 (yükseltilmiş)", "512GB", G16_SCREEN, "Windows 11 Pro", None, "TL",
     "Teknosa", "TR", "Stokta yok", "", "Ekim 2026",
     "https://www.teknosa.com/asus-rog-strix-g16-g614jvn3159a14-intel-core-i7-13650hx-16-64-gb-ram-512-gb-ssd-rtx-4060-w11p-gaming-laptop-p-785357927"),
    ("Sıfır", "ROG Strix G16 (2023)", "G614JV-N3211", "16GB DDR5", "512GB", G16_SCREEN, "-", None, "TL",
     "ASUS Türkiye resmi mağaza", "TR", "Fiyat görülemedi", "", "Ekim 2026",
     "https://shop.asus.com/tr/90nr0c61-m016j0-rog-strix-g16-2023.html"),
    # ---------------- İKİNCİ EL ----------------
    ("İkinci el", "ROG Strix G16 (2023)", "G614JV", "16GB DDR5", "-", '16" 165Hz', "-", 57000, "TL",
     "DonanımHaber forum", "TR", "İlan (durumu bilinmiyor)",
     "2,5 ay kullanılmış; hafif oyun ve video izleme; kutu, fatura ve garanti mevcut. Satıcıya göre sıfır fiyatı o tarihte 64-66 bin TL.",
     "29.11.2025", "https://forum.donanimhaber.com/asus-rog-strix-g16-i7-13650hx-rtx-4060-gaming-laptop--158114501"),
    ("İkinci el", "ROG Strix G16 (2023)", "G614JV", "-", "-", '16"', "-", 40000, "TL",
     "Technopat forum", "TR", "İlan (yaklaşık 4 gün önce)",
     "2023 model; kutusu ve faturası var; garantisi bitmiş.",
     "Ekim 2026", "https://www.technopat.net/sosyal/konu/ikinci-el-asus-rog-strix-g16-4060-laptop-alinir-mi.4187123/"),
    ("İkinci el", "ROG Strix G16", "G614JV (doğrulanmalı)", "16GB DDR5", "-", '16"', "-", 55000, "TL",
     "Technopat forum (Satılık)", "TR", "İlan",
     "'Sıfırdan farksız' deniyor. DİKKAT: ilan başlığında RTX 4070 140W, metinde RTX 4060 yazıyor. Ekran kartını satıcıya teyit ettirin.",
     "-", "https://www.technopat.net/sosyal/konu/asus-rog-strix-g16-i7-13650hx-rtx-4070-140w-16gb-ram.4116714/"),
    ("İkinci el", "ROG Strix G16", "G614JV", "24GB DDR5", "-", '16" FHD 165Hz', "-", 40000, "TL",
     "DonanımArşivi forum", "TR", "Eski ilan (~2,5 yıl önce)",
     "Kondisyon belirtilmemiş. Forumda RAM hız oranı ve 16 inç ekran zayıf yön olarak eleştirilmiş.",
     "~2024", "https://forum.donanimarsivi.com/konu/asus-rog-strix-i7-13650hx-rtx-4060-ikinci-el-40-bin-alinir-mi.898488/"),
]

FILL = {
    "Sıfır": "E2EFDA",
    "İkinci el": "FCE4D6",
    "Yenilenmiş (ikinci el)": "FFF2CC",
    "Yenilenmiş / açık kutu": "FFF2CC",
}


def records():
    for i, r in enumerate(ROWS, 1):
        durum, seri, kod, ram, ssd, ekran, os_, fiyat, pb, satici, ulke, stok, kond, tarih, link = r
        yield [i, durum, seri, kod, CPU, GPU, ram, ssd, ekran, os_, fiyat, pb,
               satici, stok, kond, tarih, link]


def write_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "İlanlar"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rec in records():
        ws.append(rec)
        row = ws.max_row
        color = FILL.get(rec[1], "FFFFFF")
        for c in ws[row]:
            c.fill = PatternFill("solid", fgColor=color)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row, 11).number_format = "#,##0.00"
        link = ws.cell(row, 17)
        link.hyperlink = link.value
        link.font = Font(color="0563C1", underline="single")
    widths = [5, 12, 22, 24, 26, 13, 20, 16, 24, 15, 12, 8, 28, 22, 60, 12, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    notes = wb.create_sheet("Notlar")
    for line in [
        "Asus ROG Strix + RTX 4060 + Intel Core i7 13. nesil (i7-13650HX) ilan taraması - SADECE TÜRKİYE",
        "Tarama tarihi: 12 Ekim 2026",
        "",
        "Bu kombinasyon 2 seride var: ROG Strix G16 (G614JV, 16\") ve ROG Strix G18 (G814JV, 18\").",
        "Hepsinde işlemci i7-13650HX (14 çekirdek), ekran kartı RTX 4060 8GB GDDR6 (TGP 140W; bazı ilanlar 175W yazıyor).",
        "2024 serisi G614JVR / G814JVR i7-14650HX (14. nesil) kullandığı için tabloya alınmadı.",
        "Sadece anakart satan 'parça' ilanları da tabloya alınmadı.",
        "",
        "Renkler: yeşil = sıfır, turuncu = ikinci el.",
        "ROG Strix G18 (G814JV) için Türkiye'de güncel TL ilanı bulunamadı.",
        "'(yükseltilmiş)' = RAM/SSD fabrika çıkışı değil, satıcı tarafından artırılmış.",
        "Boş fiyat = arama sonucunda fiyat görünmedi; linkten kontrol edin.",
        "",
        "SINIRLAMA: Ağ politikası nedeniyle ilan sayfaları doğrudan açılamadı. Veriler arama motoru",
        "özetlerinden alındı. Fiyat/stok anlık görüntüdür ve bazıları eskidir (Veri Tarihi sütunu).",
        "sahibinden.com, letgo ve dolap ilanları aramada çıkmadı; bu siteleri ayrıca kontrol edin.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    notes["A1"].font = Font(bold=True, size=13)
    wb.save(path)


def write_csv(path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(records())


if __name__ == "__main__":
    write_xlsx(OUT_DIR / f"{BASENAME}.xlsx")
    write_csv(OUT_DIR / f"{BASENAME}.csv")
    print(f"{len(ROWS)} satır yazıldı")
